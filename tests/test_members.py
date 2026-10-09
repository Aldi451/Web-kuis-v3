"""
Tes fitur Member (peserta terdaftar) & roster kuis:

  * Akun Member dibuat (Admin/Host), login mengembalikan token sesi.
  * Halaman awal = login only; role Member diarahkan ke portal member (dijaga test static assets).
  * Host memilih member yang ikut kuis + levelnya (roster). Validasi roster & pool soal per level.
  * Member login lalu join (scan barcode): soal diacak dari level roster, terhubung ke akunnya.
  * Roster ditegakkan: hanya member terdaftar yang bisa gabung; member boleh menyusul saat kuis berjalan.
  * Anti-nyontek: soal yang pernah dikerjakan member tidak diundi lagi (lintas room, juga setelah leave/rejoin).
  * Import soal dari Excel (.xlsx) + download template.
"""
import io
import uuid

import pytest
from fastapi.testclient import TestClient

import app as quiz_app
import database

LEVELS = ("easy", "normal", "hard")


@pytest.fixture(scope="module")
def client():
    with TestClient(quiz_app.app) as c:
        yield c


def uniq(prefix="x"):
    return f"{prefix}-{uuid.uuid4().hex[:8]}"


def make_question(client, level="normal", text=None):
    r = client.post("/api/questions", json={
        "question_text": text or f"Soal {level} {uuid.uuid4().hex[:6]}",
        "option_a": "a", "option_b": "b", "option_c": "c", "option_d": "d",
        "correct_option": "A", "category": "Tes member", "question_type": "mcq", "level": level,
    })
    assert r.status_code == 200, r.text
    return r.json()


def make_member(client, name=None):
    name = name or uniq("member")
    r = client.post("/api/auth/register", json={"username": name, "password": "pw123", "role": "Member"})
    assert r.status_code == 200 and r.json()["role"] == "Member", r.text
    return r.json()


def login(client, username, password="pw123"):
    r = client.post("/api/auth/login", json={"username": username, "password": password})
    assert r.status_code == 200, r.text
    return r.json()


def create_room(client, ids, members=None, counts=None, mode="random", expect=200):
    body = {"title": "Kuis Member", "duration": 10, "passing_grade": 50, "question_ids": ids,
            "created_by": "admin", "question_mode": mode}
    if counts is not None:
        body["level_counts"] = counts
    if members is not None:
        body["members"] = members
    r = client.post("/api/rooms", json=body)
    assert r.status_code == expect, r.text
    return r.json()


def join(client, code, name, auth_token=None, token=None, expect=200):
    headers = {"X-Auth-Token": auth_token} if auth_token else {}
    body = {"room_id": code, "name": name, "department": "QA"}
    if token:
        body["token"] = token
    r = client.post("/api/rooms/join", json=body, headers=headers)
    assert r.status_code == expect, r.text
    return r.json()


def question_levels(client, ids):
    qs = {q["id"]: q for q in client.get("/api/questions").json()}
    return {qid: qs[qid]["level"] for qid in ids}


def start(client, code):
    r = client.put(f"/api/rooms/{code}/status", json={"status": "On Progress"})
    assert r.status_code == 200, r.text


# ───────────── akun & sesi login Member ─────────────

def test_member_register_login_and_session_token(client):
    member = make_member(client)
    assert member["role"] == "Member"

    logged = login(client, member["username"])
    assert logged["role"] == "Member" and logged["token"]
    assert database.get_user_by_token(logged["token"])["id"] == member["id"]
    # token tidak dikenal -> None
    assert database.get_user_by_token("token-ngawur") is None

    # daftar user bisa difilter per role (dipakai Host saat menyusun roster)
    members = client.get("/api/auth/users?role=Member").json()
    assert all(u["role"] == "Member" for u in members)
    assert any(u["username"] == member["username"] for u in members)
    hosts = client.get("/api/auth/users?role=Host").json()
    assert all(u["role"] == "Host" for u in hosts)


def test_delete_user_cleans_sessions_and_roster(client):
    member = make_member(client)
    token = login(client, member["username"])["token"]
    q = make_question(client)
    room = create_room(client, [q["id"]], members=[{"user_id": member["id"], "level": "normal"}], mode="fixed")
    assert any(m["user_id"] == member["id"] for m in client.get(f"/api/rooms/{room['room_code']}/members").json())

    assert client.delete(f"/api/auth/users/{member['id']}").status_code == 200
    assert database.get_user_by_token(token) is None                       # sesi login ikut terhapus
    assert client.get(f"/api/rooms/{room['room_code']}/members").json() == []  # roster ikut terhapus


# ───────────── roster: validasi saat membuat kuis ─────────────

def test_room_rejects_roster_with_non_member_or_bad_level(client):
    host = client.post("/api/auth/register", json={"username": uniq("host"), "password": "pw", "role": "Host"}).json()
    q = make_question(client)
    # user_id milik Host (bukan Member) -> ditolak
    create_room(client, [q["id"]], members=[{"user_id": host["id"], "level": "easy"}], mode="fixed", expect=400)
    # user_id tidak ada -> ditolak
    create_room(client, [q["id"]], members=[{"user_id": 99999999, "level": "easy"}], mode="fixed", expect=400)
    # level tidak valid -> 422 (validasi pydantic)
    member = make_member(client)
    r = client.post("/api/rooms", json={
        "title": "x", "duration": 5, "passing_grade": 50, "question_ids": [q["id"]],
        "members": [{"user_id": member["id"], "level": "super"}]})
    assert r.status_code == 422


def test_room_random_validates_pool_against_roster_levels(client):
    member = make_member(client)
    easy = [make_question(client, "easy")["id"] for _ in range(2)]
    # roster minta 3 soal easy per member, tetapi pool hanya punya 2 -> ditolak
    room = create_room(client, easy, members=[{"user_id": member["id"], "level": "easy"}],
                       counts={"easy": 3, "normal": 0, "hard": 0}, expect=400)
    # pool cukup -> berhasil
    easy += [make_question(client, "easy")["id"]]
    room = create_room(client, easy, members=[{"user_id": member["id"], "level": "easy"}],
                       counts={"easy": 3, "normal": 0, "hard": 0})
    assert room["room_code"]


def test_room_members_get_and_update(client):
    m1, m2 = make_member(client), make_member(client)
    q = make_question(client)
    room = create_room(client, [q["id"]], members=[{"user_id": m1["id"], "level": "easy"}], mode="fixed")
    code = room["room_code"]

    members = client.get(f"/api/rooms/{code}/members").json()
    assert [(m["user_id"], m["level"]) for m in members] == [(m1["id"], "easy")]
    assert members[0]["username"] == m1["username"]

    # Ganti roster: m1 naik level, m2 masuk
    r = client.put(f"/api/rooms/{code}/members", json={
        "members": [{"user_id": m1["id"], "level": "hard"}, {"user_id": m2["id"], "level": "normal"}]})
    assert r.status_code == 200, r.text
    members = r.json()["members"]
    assert {(m["user_id"], m["level"]) for m in members} == {(m1["id"], "hard"), (m2["id"], "normal")}

    # Setelah kuis dimulai, roster tidak bisa diubah lagi
    start(client, code)
    r = client.put(f"/api/rooms/{code}/members", json={"members": []})
    assert r.status_code == 409


# ───────────── join: roster ditegakkan & level dipakai ─────────────

def test_member_join_links_account_and_uses_roster_level(client):
    member = make_member(client)
    token = login(client, member["username"])["token"]
    easy = [make_question(client, "easy") for _ in range(3)]
    hard = [make_question(client, "hard") for _ in range(3)]
    ids = [q["id"] for q in easy + hard]
    room = create_room(client, ids, members=[{"user_id": member["id"], "level": "hard"}],
                       counts={"easy": 0, "normal": 0, "hard": 2})

    joined = join(client, room["room_code"], member["username"], auth_token=token)
    assert joined["assigned_level"] == "hard"
    assert joined["is_member"] is True
    assert len(joined["question_ids"]) == 2
    # semua soal miliknya berlevel hard (level roster dipakai, bukan campuran)
    assert set(question_levels(client, joined["question_ids"]).values()) == {"hard"}

    # Data peserta di server terhubung ke akun member
    participants = client.get(f"/api/rooms/{room['room_code']}/participants").json()
    assert participants[0]["assigned_level"] == "hard" and participants[0]["is_member"] is True


def test_roster_enforced_for_join(client):
    m1, m2 = make_member(client), make_member(client)
    t1, t2 = login(client, m1["username"])["token"], login(client, m2["username"])["token"]
    q = make_question(client)
    room = create_room(client, [q["id"]], members=[{"user_id": m1["id"], "level": "normal"}], mode="fixed")
    code = room["room_code"]

    # Member yang TIDAK terdaftar di roster -> ditolak
    join(client, code, m2["username"], auth_token=t2, expect=400)
    # Tanpa login (anonymous) -> ditolak karena room punya roster
    join(client, code, "Tanpa Login", expect=400)
    # Host/Admin juga tidak bisa join lewat akun non-member
    admin_token = login(client, "admin", "admin123")["token"]
    join(client, code, "admin", auth_token=admin_token, expect=400)
    # Member terdaftar -> berhasil
    join(client, code, m1["username"], auth_token=t1)


def test_rostered_member_can_join_after_quiz_started(client):
    m1, m2 = make_member(client), make_member(client)
    t1, t2 = login(client, m1["username"])["token"], login(client, m2["username"])["token"]
    hard = [make_question(client, "hard") for _ in range(2)]
    room = create_room(client, [q["id"] for q in hard],
                       members=[{"user_id": m1["id"], "level": "hard"}, {"user_id": m2["id"], "level": "hard"}],
                       counts={"easy": 0, "normal": 0, "hard": 2})
    code = room["room_code"]

    join(client, code, m1["username"], auth_token=t1)
    start(client, code)

    # Member roster boleh scan & gabung MESKI kuis sudah berjalan (soal langsung diundi)
    late = join(client, code, m2["username"], auth_token=t2)
    assert late["assigned_level"] == "hard" and len(late["question_ids"]) == 2
    # Peserta tanpa akun tetap tidak bisa menyusul saat kuis berjalan
    join(client, code, "Orang Baru", expect=400)


def test_member_without_roster_joins_normally(client):
    member = make_member(client)
    token = login(client, member["username"])["token"]
    ids = [make_question(client, "easy")["id"], make_question(client, "hard")["id"]]
    room = create_room(client, ids, counts={"easy": 1, "normal": 0, "hard": 1})
    joined = join(client, room["room_code"], member["username"], auth_token=token)
    assert joined["assigned_level"] is None
    assert len(joined["question_ids"]) == 2


def test_member_resume_by_account_after_browser_cleared(client):
    """HP member di-restart (localStorage lama hilang): login lagi -> sesi kuisnya tetap bisa dilanjutkan."""
    member = make_member(client)
    token = login(client, member["username"])["token"]
    ids = [make_question(client, "normal")["id"] for _ in range(2)]
    room = create_room(client, ids, members=[{"user_id": member["id"], "level": "normal"}],
                       counts={"easy": 0, "normal": 1, "hard": 0})
    code = room["room_code"]

    first = join(client, code, member["username"], auth_token=token)
    # "localStorage hilang": join lagi HANYA dengan token login (tanpa token peserta lama)
    resumed = join(client, code, member["username"], auth_token=token)
    assert resumed["resumed"] is True
    assert resumed["id"] == first["id"] and resumed["token"] == first["token"]
    assert resumed["question_ids"] == first["question_ids"]


# ───────────── anti-nyontek: soal tidak berulang ─────────────

def test_member_never_gets_a_question_they_already_did(client):
    member = make_member(client)
    token = login(client, member["username"])["token"]
    easy = [make_question(client, "easy") for _ in range(4)]
    ids = [q["id"] for q in easy]

    room1 = create_room(client, ids, members=[{"user_id": member["id"], "level": "easy"}],
                        counts={"easy": 2, "normal": 0, "hard": 0})
    first = join(client, room1["room_code"], member["username"], auth_token=token)
    assert len(first["question_ids"]) == 2

    # Kuis berikutnya (room baru, pool sama): soal yang pernah dikerjakan TIDAK muncul lagi
    room2 = create_room(client, ids, members=[{"user_id": member["id"], "level": "easy"}],
                        counts={"easy": 2, "normal": 0, "hard": 0})
    second = join(client, room2["room_code"], member["username"], auth_token=token)
    assert set(second["question_ids"]).isdisjoint(first["question_ids"])
    assert len(second["question_ids"]) == 2


def test_leave_and_rejoin_does_not_reset_seen_questions(client):
    """Celah: keluar lalu gabung lagi tidak boleh mengundi soal yang sudah pernah dilihat."""
    member = make_member(client)
    token = login(client, member["username"])["token"]
    ids = [make_question(client, "easy")["id"] for _ in range(4)]
    room = create_room(client, ids, members=[{"user_id": member["id"], "level": "easy"}],
                       counts={"easy": 2, "normal": 0, "hard": 0})
    code = room["room_code"]

    first = join(client, code, member["username"], auth_token=token)
    # Keluar dari waiting room (menghapus participant & assignment-nya)
    r = client.post(f"/api/participants/{first['id']}/leave", headers={"X-Participant-Token": first["token"]})
    assert r.status_code == 200, r.text
    # Gabung lagi: soal tetap mengecualikan yang pernah dilihat (anti-nyontek)
    again = join(client, code, member["username"], auth_token=token)
    assert set(again["question_ids"]).isdisjoint(first["question_ids"])


def test_two_members_same_room_get_different_question_sets(client):
    m1, m2 = make_member(client), make_member(client)
    t1, t2 = login(client, m1["username"])["token"], login(client, m2["username"])["token"]
    ids = [make_question(client, "hard")["id"] for _ in range(4)]
    room = create_room(client, ids,
                       members=[{"user_id": m1["id"], "level": "hard"}, {"user_id": m2["id"], "level": "hard"}],
                       counts={"easy": 0, "normal": 0, "hard": 2})
    j1 = join(client, room["room_code"], m1["username"], auth_token=t1)
    j2 = join(client, room["room_code"], m2["username"], auth_token=t2)
    assert set(j1["question_ids"]).isdisjoint(j2["question_ids"])


# ───────────── import soal dari Excel + template ─────────────

def _xlsx_bytes(rows):
    openpyxl = pytest.importorskip("openpyxl")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Soal"
    ws.append(["Pertanyaan", "Pilihan A", "Pilihan B", "Pilihan C", "Pilihan D",
               "Kunci Jawaban (A/B/C/D atau teks essay)", "Kategori", "Tipe (mcq/essay)", "Level (easy/normal/hard)"])
    for row in rows:
        ws.append(row)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def test_import_template_download(client):
    pytest.importorskip("openpyxl")
    from openpyxl import load_workbook
    r = client.get("/api/questions/import/template")
    assert r.status_code == 200
    assert "spreadsheetml" in r.headers["content-type"]
    assert "template_import_soal.xlsx" in r.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(r.content))
    assert "Soal" in wb.sheetnames and "Petunjuk" in wb.sheetnames
    header = [cell.value for cell in wb["Soal"][1]]
    assert "Pertanyaan" in header and "Kunci Jawaban (A/B/C/D atau teks essay)" in header


def test_import_questions_from_excel(client):
    before = len(client.get("/api/questions").json())
    data = _xlsx_bytes([
        ["Berapa 3+3?", "5", "6", "7", "8", "B", "Umum", "mcq", "easy"],
        ["Jelaskan K3!", "", "", "", "", "Utamakan keselamatan kerja", "K3", "essay", "normal"],
        ["Soal rusak", "a", "", "c", "d", "A", "Umum", "mcq", "easy"],  # pilihan B kosong -> error
    ])
    r = client.post("/api/questions/import", files={
        "file": ("soal.xlsx", data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert r.status_code == 200, r.text
    result = r.json()
    assert result["imported"] == 2 and result["skipped"] == 1
    assert result["errors"][0]["row"] == 4 and "Pilihan" in result["errors"][0]["message"]

    questions = client.get("/api/questions").json()
    assert len(questions) == before + 2
    by_text = {q["question_text"]: q for q in questions}
    assert by_text["Berapa 3+3?"]["correct_option"] == "B" and by_text["Berapa 3+3?"]["level"] == "easy"
    assert by_text["Jelaskan K3!"]["question_type"] == "essay" and by_text["Jelaskan K3!"]["level"] == "normal"


def test_import_rejects_non_excel_file(client):
    r = client.post("/api/questions/import", files={"file": ("soal.xlsx", b"bukan excel", "application/octet-stream")})
    assert r.status_code == 400
