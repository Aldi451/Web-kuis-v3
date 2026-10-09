"""
excel_import.py - impor bank soal dari file Excel (.xlsx) + pembuatan template Excel.

Dipakai endpoint:
  * GET  /api/questions/import/template  -> unduh template .xlsx (Format & contoh pengisian jelas)
  * POST /api/questions/import           -> unggah .xlsx, soal valid langsung masuk bank soal

Aturan pengisian (kolom, urutan tidak harus sama - yang penting nama kolomnya cocok):
  question_text  : teks pertanyaan (wajib)
  option_a..d    : pilihan jawaban A-D (wajib untuk soal pilihan ganda)
  correct_answer : kunci jawaban. Pilihan ganda: huruf A/B/C/D. Essay: teks jawaban referensi.
  category       : kategori soal (kosong = "Umum")
  question_type  : "mcq" (pilihan ganda, bawaan) atau "essay"
  level          : easy / normal / hard (kosong = normal). Alias Indonesia: mudah/sedang/sulit.

Butuh library openpyxl (ada di requirements.txt). Tanpa openpyxl, endpoint mengembalikan 501.
"""

import io
import re

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

# Urutan & label kolom di template (pencocokan saat import berdasarkan NAMA kolom, bukan posisi)
COLUMNS = [
    ("question_text", "Pertanyaan", 50),
    ("option_a", "Pilihan A", 30),
    ("option_b", "Pilihan B", 30),
    ("option_c", "Pilihan C", 30),
    ("option_d", "Pilihan D", 30),
    ("correct_answer", "Kunci Jawaban (A/B/C/D atau teks essay)", 24),
    ("category", "Kategori", 16),
    ("question_type", "Tipe (mcq/essay)", 14),
    ("level", "Level (easy/normal/hard)", 18),
]

LEVEL_ALIASES = {
    "easy": "easy", "mudah": "easy",
    "normal": "normal", "sedang": "normal", "medium": "normal",
    "hard": "hard", "sulit": "hard",
}

VALID_TYPES = ("mcq", "essay")

MAX_IMPORT_ROWS = 2000  # batas wajar agar satu file tidak mengunci server terlalu lama

TEMPLATE_FILENAME = "template_import_soal.xlsx"


def _require_openpyxl():
    if not HAS_OPENPYXL:
        raise RuntimeError("Library openpyxl belum terpasang. Jalankan: pip install openpyxl")


def build_template_bytes() -> bytes:
    """Buat file template .xlsx: sheet 'Soal' (header + contoh) + sheet 'Petunjuk'."""
    _require_openpyxl()
    wb = Workbook()

    # ── Sheet 1: Soal ──
    ws = wb.active
    ws.title = "Soal"
    header_fill = PatternFill("solid", fgColor="283E94")
    header_font = Font(bold=True, color="FFFFFF")
    for col_index, (key, label, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_index, value=label)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col_index)].width = width
    ws.row_dimensions[1].height = 30

    example_rows = [
        ["Berapa hasil 2 + 2?", "3", "4", "5", "6", "B", "Umum", "mcq", "easy"],
        ["Jelaskan langkah check-in tamu!", "", "", "", "",
         "Tamu mengisi buku tamu, menunjukkan KTP, resepsioner input data ke sistem, lalu serahkan kunci kamar.",
         "Front Office", "essay", "normal"],
    ]
    for row_index, row in enumerate(example_rows, start=2):
        for col_index, value in enumerate(row, start=1):
            cell = ws.cell(row=row_index, column=col_index, value=value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"

    # ── Sheet 2: Petunjuk ──
    guide = wb.create_sheet("Petunjuk")
    guide.column_dimensions["A"].width = 110
    lines = [
        "CARA IMPORT SOAL DARI EXCEL",
        "",
        "1. Isi soal pada sheet 'Soal', mulai dari baris 2 (baris 1 = judul kolom, jangan dihapus).",
        "2. Nama kolom harus sama seperti di template (urutan kolom boleh berbeda).",
        "3. Kolom 'Pertanyaan' wajib diisi. Baris yang teks pertanyaannya kosong dilewati.",
        "4. Tipe 'mcq' (pilihan ganda): isi Pilihan A-D dan Kunci Jawaban dengan huruf A/B/C/D.",
        "5. Tipe 'essay': cukup isi Kunci Jawaban dengan teks jawaban referensi (pilihan A-D boleh kosong).",
        "6. Level: easy (mudah), normal (sedang), atau hard (sulit). Kosong = normal.",
        "7. Simpan file (tetap .xlsx), lalu upload lewat tombol 'Import Excel' di Bank Soal.",
        "8. Baris yang bermasalah dilaporkan (nomor baris + alasannya) dan TIDAK di-import;",
        "   perbaiki di Excel lalu import lagi. Perhatikan: soal yang SAMA persis bisa ter-import dua kali,",
        "   jadi cek dahulu Bank Soal (atau hapus soal lama) sebelum import file yang sama.",
        "",
        f"Batas maksimal {MAX_IMPORT_ROWS} soal per sekali import.",
    ]
    for index, text in enumerate(lines, start=1):
        cell = guide.cell(row=index, column=1, value=text)
        if index == 1:
            cell.font = Font(bold=True, size=13)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _canon(value) -> str:
    """'Kunci Jawaban (A/B/C/D atau teks essay)' -> 'kunci_jawaban_a_b_c_d_atau_teks_essay'."""
    return re.sub(r"[^a-z0-9]+", "_", str(value or "").strip().lower()).strip("_")


# Header boleh pakai nama kolom Inggris (question_text) maupun label Indonesia dari template (Pertanyaan)
COLUMN_ALIASES = {
    "question_text": {"question_text", "pertanyaan", "soal", "question", "teks_pertanyaan"},
    "option_a": {"option_a", "pilihan_a", "jawaban_a", "a"},
    "option_b": {"option_b", "pilihan_b", "jawaban_b", "b"},
    "option_c": {"option_c", "pilihan_c", "jawaban_c", "c"},
    "option_d": {"option_d", "pilihan_d", "jawaban_d", "d"},
    "correct_answer": {"correct_answer", "kunci_jawaban", "kunci", "jawaban_benar", "jawaban", "answer",
                       "kunci_jawaban_a_b_c_d_atau_teks_essay", "kunci_jawaban_a_b_c_d", "kunci_jawaban_essay"},
    "category": {"category", "kategori"},
    "question_type": {"question_type", "tipe", "tipe_soal", "type", "jenis_soal", "jenis", "tipe_mcq_essay"},
    "level": {"level", "tingkat", "level_soal", "tingkat_kesulitan", "level_easy_normal_hard"},
}


def _header_to_column(value) -> str:
    canon = _canon(value)
    for column, aliases in COLUMN_ALIASES.items():
        if canon in aliases:
            return column
    return ""


def parse_questions_excel(data: bytes):
    """
    Parse file .xlsx -> (rows, errors).
    rows   : list[dict] siap masuk tabel questions (sudah divalidasi & dinormalisasi).
    errors : list[{"row": int, "message": str}] untuk baris yang bermasalah (1-based, baris 1 = header).
    Raise ValueError jika file bukan .xlsx yang valid / tidak ada sheet.
    """
    _require_openpyxl()
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError(f"File bukan format Excel (.xlsx) yang valid: {exc}")

    ws = wb["Soal"] if "Soal" in wb.sheetnames else wb.worksheets[0]
    rows_iter = ws.iter_rows(values_only=True)

    # Baris 1 = header: cocokkan nama kolom (urutan bebas)
    try:
        header = next(rows_iter)
    except StopIteration:
        wb.close()
        raise ValueError("Sheet kosong. Isi judul kolom sesuai template (sheet 'Petunjuk' di template).")
    column_index = {}
    for position, cell in enumerate(header):
        column = _header_to_column(cell)
        if column and column not in column_index:
            column_index[column] = position
    if "question_text" not in column_index:
        wb.close()
        raise ValueError("Kolom 'Pertanyaan' (question_text) tidak ditemukan. Unduh template untuk melihat format yang benar.")

    rows, errors = [], []
    for row_number, raw in enumerate(rows_iter, start=2):
        if row_number > MAX_IMPORT_ROWS + 1:
            errors.append({"row": row_number, "message": f"Melebihi batas maksimal {MAX_IMPORT_ROWS} soal per import."})
            break
        values = list(raw) + [None] * len(COLUMNS)  # baris pendek tetap aman diindeks

        def cell(key):
            index = column_index.get(key)
            return values[index] if index is not None and index < len(values) else None

        question_text = str(cell("question_text") or "").strip()
        if not question_text:
            continue  # baris kosong / tanpa pertanyaan -> dilewati diam-diam

        row_data = {"question_text": question_text}
        error = None

        raw_type = str(cell("question_type") or "mcq").strip().lower()
        question_type = raw_type if raw_type in VALID_TYPES else None
        if question_type is None:
            error = f"Tipe soal '{raw_type}' tidak dikenal. Gunakan: mcq atau essay."
        row_data["question_type"] = question_type or "mcq"

        category = str(cell("category") or "").strip() or "Umum"
        row_data["category"] = category[:100]

        raw_level = str(cell("level") or "").strip().lower()
        level = LEVEL_ALIASES.get(raw_level, "normal") if raw_level else "normal"
        if raw_level and raw_level not in LEVEL_ALIASES:
            error = (error + " " if error else "") + \
                f"Level '{raw_level}' tidak dikenal. Gunakan: easy, normal, atau hard."
        row_data["level"] = level

        if row_data["question_type"] == "mcq":
            options = [str(cell(f"option_{letter}") or "").strip() for letter in ("a", "b", "c", "d")]
            if not all(options):
                error = (error + " " if error else "") + \
                    "Pilihan A, B, C, dan D wajib diisi untuk soal pilihan ganda."
            row_data["option_a"], row_data["option_b"], row_data["option_c"], row_data["option_d"] = options
            correct = str(cell("correct_answer") or "").strip().upper()
            if correct not in ("A", "B", "C", "D"):
                error = (error + " " if error else "") + \
                    "Kunci jawaban pilihan ganda harus berupa huruf A, B, C, atau D."
            row_data["correct_answer"] = correct
        else:  # essay
            row_data["option_a"] = row_data["option_b"] = row_data["option_c"] = row_data["option_d"] = ""
            correct = str(cell("correct_answer") or "").strip()
            if not correct:
                error = (error + " " if error else "") + \
                    "Kunci jawaban essay wajib diisi (teks jawaban referensi)."
            row_data["correct_answer"] = correct

        if error:
            errors.append({"row": row_number, "message": error})
        else:
            rows.append(row_data)

    wb.close()
    return rows, errors
